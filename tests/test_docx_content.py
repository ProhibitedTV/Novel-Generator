from docx import Document

from novel_generator.services.docx_content import add_manuscript_content


def render(text, tmp_path):
    document = Document()
    add_manuscript_content(document, text)
    path = tmp_path / "book.docx"
    document.save(path)
    return Document(path)


def test_handbook_headings_nested_emphasis_and_quote_lists_survive_word_roundtrip(tmp_path):
    document = render("### Record fields\n\nThe **Item ID** is *unique*.\n\n> **Example**\n>\n> * **ID:** HYP-1\n> * Status: available", tmp_path)
    assert document.paragraphs[0].text == "Record fields"
    assert document.paragraphs[0].style.name == "Heading 4"
    body = document.paragraphs[1]
    assert any(run.text == "Item ID" and run.bold for run in body.runs)
    assert any(run.text == "unique" and run.italic for run in body.runs)
    assert [p.text for p in document.paragraphs[-2:]] == ["\u2022 ID: HYP-1", "\u2022 Status: available"]


def test_poetry_and_speaker_lines_keep_breaks_and_literal_symbols(tmp_path):
    text = "**ADA:** Open the gate.\n**BEN:** After you.\n\nThe river bends,  \nthen rests.\nA_B_1 and 2 * 3 remain."
    document = render(text, tmp_path)
    assert document.paragraphs[0].text == "ADA: Open the gate.\nBEN: After you."
    assert document.paragraphs[1].text == "The river bends,\nthen rests.\nA_B_1 and 2 * 3 remain."


def test_numbered_lists_keep_start_values_restart_and_nested_bullets(tmp_path):
    document = render("7. Verify status.\n8. Inspect condition.\n   - Record differences.\n\nInterlude.\n\n0. Start a separate procedure.\n1. Continue.", tmp_path)
    assert [p.text for p in document.paragraphs] == [
        "7. Verify status.", "8. Inspect condition.", "\u2022 Record differences.", "Interlude.",
        "0. Start a separate procedure.", "1. Continue."]


def test_reference_tables_preserve_fields_and_numbers(tmp_path):
    document = render("| Field | Value |\n| --- | --- |\n| **Volume** | 10 mL |\n| ID | A_B_1 |", tmp_path)
    assert len(document.tables) == 1
    assert [[cell.text for cell in row.cells] for row in document.tables[0].rows] == [
        ["Field", "Value"], ["Volume", "10 mL"], ["ID", "A_B_1"]]
    assert document.tables[0].cell(0, 0).paragraphs[0].runs[0].bold


def test_code_and_source_links_remain_readable_without_fetching_assets(tmp_path):
    document = render("See [Desk Procedure](https://example.org/manual).\n\n```python\na = '**literal**'\n```\n\nUse `A_B_1`.\n\n<script>literal text</script>", tmp_path)
    assert [p.text for p in document.paragraphs] == [
        "See Desk Procedure (https://example.org/manual).", "a = '**literal**'\n", "Use A_B_1.",
        "<script>literal text</script>"]
    assert document.paragraphs[1].runs[0].font.name == "Courier New"
