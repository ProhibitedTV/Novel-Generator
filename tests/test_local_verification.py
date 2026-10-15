from types import SimpleNamespace
import json
import pytest
from datetime import datetime

from novel_generator.models import ChapterStatus, RunStatus
from novel_generator.services.local_verification import verification_report
from novel_generator.services.autonomous_editorial import _hash
from novel_generator.services.local_verification import load_project_spec
from novel_generator.services.exports import render_markdown, render_docx


def _run(tmp_path):
    artifacts = []
    for kind in ("markdown", "docx", "qa-report"):
        (tmp_path / kind).write_text("test artifact", encoding="utf-8")
        artifacts.append(SimpleNamespace(kind=kind, relative_path=kind))
    run = SimpleNamespace(
        id="test-run", model_name="local-model", status=RunStatus.COMPLETED, error_message=None,
        project=SimpleNamespace(title='Synthetic test book',story_brief={}),
        requested_chapters=2, target_word_count=1000, artifacts=artifacts,
        events=[SimpleNamespace(event_type="summary_fallback", payload={"chapter_number": 1}, sequence=1, created_at=datetime.utcnow())],
        stage_attempts=[],
        chapters=[SimpleNamespace(chapter_number=number, title=f'Turn {number}', status=ChapterStatus.COMPLETED,
                                  content=f"Some actual prose for section {number}.", summary="Summary", continuity_update={"world_state": "Known"},
                                  qa_notes={}) for number in (1, 2)],
    )
    (tmp_path/'markdown').write_text(render_markdown(run.project,run.chapters),encoding='utf-8')
    render_docx(run.project,run.chapters,tmp_path/'docx')
    return run


def _event(run, kind, payload):
    run.events.append(SimpleNamespace(event_type=kind, payload=payload,
        sequence=len(run.events) + 1, created_at=datetime.utcnow()))


def _accept(run):
    hashes = {str(c.chapter_number): _hash(c.content) for c in run.chapters}
    def review(scope, numbers, context=None):
        _event(run, 'autonomous_review_completed', {'scope': scope,
            'source_hashes': {str(n): hashes[str(n)] for n in numbers},
            'context_source_hashes': context or {},
            'review': {'reviewed_chapters': numbers, 'plot_coherent': True, 'canon_consistent': True,
                'prose_clean': True, 'ending_complete': True, 'issues': []}})
    for number in (1, 2):
        for scope in ('final_chapter', 'canon_identity', 'developmental_chapter'):
            review(scope, [number], hashes if scope == 'developmental_chapter' else None)
    review('adjacent_chapters', [1, 2])
    review('whole_book', [1, 2])
    _event(run, 'autonomous_quality_passed', {
        'manuscript_hash': _hash([(c.chapter_number, c.content) for c in run.chapters]), 'acceptance_contract': 2})


def test_completion_report_retains_quality_and_length_limits(tmp_path):
    report = verification_report(_run(tmp_path), tmp_path)
    assert report["pipeline_complete"]
    assert report["actual_words"] == 12
    assert report["target_words"] == 1000
    assert report["fallback_events"][0]["type"] == "summary_fallback"
    assert "not proof of logical consistency" in report["quality_note"]


def test_completion_report_rejects_missing_checkpoint_or_chapter(tmp_path):
    run = _run(tmp_path)
    run.chapters[0].continuity_update = None
    assert not verification_report(run, tmp_path)["pipeline_complete"]
    run.chapters.pop(0)
    assert not verification_report(run, tmp_path)["pipeline_complete"]


def test_completion_report_requires_all_export_formats(tmp_path):
    run = _run(tmp_path)
    (tmp_path / "docx").unlink()
    assert not verification_report(run, tmp_path)["pipeline_complete"]
    run.artifacts = [artifact for artifact in run.artifacts if artifact.kind != "docx"]
    assert not verification_report(run, tmp_path)["pipeline_complete"]


def test_automatic_acceptance_must_match_current_manuscript(tmp_path):
    run = _run(tmp_path)
    _accept(run)
    assert verification_report(run, tmp_path)["autonomous_checks_passed"]
    run.chapters[0].content = "A different ending."
    assert not verification_report(run, tmp_path)["autonomous_checks_passed"]


def test_later_current_source_concerns_revoke_qualification_success(tmp_path):
    run = _run(tmp_path)
    _accept(run)
    _event(run, 'editorial_concerns_recorded', {
        'manuscript_hash': _hash([(c.chapter_number, c.content) for c in run.chapters]),
        'message': 'Synthetic source inconsistency requires revision.'})
    report = verification_report(run, tmp_path)
    assert report['pipeline_complete']
    assert not report['autonomous_checks_passed']
    assert not report['editorial_coverage']['accepted']
    assert report['current_source_concerns'] == 1


@pytest.mark.parametrize('missing_scope', ['final_chapter','canon_identity','developmental_chapter','adjacent_chapters','whole_book'])
def test_pass_event_cannot_replace_missing_current_source_audits(tmp_path,missing_scope):
    run = _run(tmp_path)
    _accept(run)
    run.events = [e for e in run.events if e.payload.get('scope') != missing_scope]
    assert not verification_report(run,tmp_path)['autonomous_checks_passed']


def test_old_concerns_for_different_source_do_not_revoke_new_acceptance(tmp_path):
    run = _run(tmp_path)
    _event(run,'editorial_concerns_recorded',{'manuscript_hash':'old-source'})
    _accept(run)
    report=verification_report(run,tmp_path)
    assert report['autonomous_checks_passed'] and report['current_source_concerns']==0


def test_legacy_acceptance_is_retained_but_cannot_qualify_current_gate(tmp_path):
    run=_run(tmp_path)
    _accept(run)
    run.events[-1].payload['acceptance_contract']=1
    run.events=[e for e in run.events if e.payload.get('scope') not in {'canon_identity','developmental_chapter'}]
    report=verification_report(run,tmp_path)
    assert report['editorial_coverage']['accepted']
    assert report['editorial_coverage']['acceptance_contract']==1
    assert not report['autonomous_checks_passed']


@pytest.mark.parametrize('format',['markdown','docx'])
def test_export_qualification_rejects_stale_or_truncated_content(tmp_path,format):
    run=_run(tmp_path)
    _accept(run)
    shortened=run.chapters[:1]
    if format=='markdown':
        (tmp_path/format).write_text(render_markdown(run.project,shortened),encoding='utf-8')
    else:
        render_docx(run.project,shortened,tmp_path/format)
    report=verification_report(run,tmp_path)
    assert report['editorial_coverage']['accepted']
    assert not report['export_integrity'][format]['matches_current_source']
    assert not report['pipeline_complete'] and not report['autonomous_checks_passed']


@pytest.mark.parametrize('format',['markdown','docx'])
def test_export_qualification_requires_current_chapter_order(tmp_path,format):
    run=_run(tmp_path)
    chapters=list(reversed(run.chapters))
    if format=='markdown':
        (tmp_path/format).write_text(render_markdown(run.project,chapters),encoding='utf-8')
    else:
        render_docx(run.project,chapters,tmp_path/format)
    assert not verification_report(run,tmp_path)['export_integrity'][format]['matches_current_source']


def test_unreadable_word_file_is_failed_evidence_instead_of_report_crash(tmp_path):
    run=_run(tmp_path)
    (tmp_path/'docx').write_bytes(b'not a Word ZIP file')
    report=verification_report(run,tmp_path)
    assert not report['pipeline_complete']
    assert report['export_integrity']['docx']['error']


def test_word_inline_formatting_and_table_cells_are_checked_without_zip_metadata(tmp_path):
    from docx import Document
    run=_run(tmp_path)
    run.chapters[0].content='A **bold** step and *quiet* reply.\n\n| Item | Count |\n| --- | --- |\n| Pears | 3 |'
    (tmp_path/'markdown').write_text(render_markdown(run.project,run.chapters),encoding='utf-8')
    render_docx(run.project,run.chapters,tmp_path/'docx')
    document=Document(tmp_path/'docx')
    document.core_properties.author='Different export metadata'
    document.save(tmp_path/'docx')
    assert verification_report(run,tmp_path)['pipeline_complete']
    document.tables[0].cell(1,1).text='4'
    document.save(tmp_path/'docx')
    assert not verification_report(run,tmp_path)['export_integrity']['docx']['matches_current_source']


def test_word_line_breaks_are_content_for_poetry_and_code(tmp_path):
    from docx import Document
    run=_run(tmp_path)
    run.chapters[0].content='First line\nSecond line'
    (tmp_path/'markdown').write_text(render_markdown(run.project,run.chapters),encoding='utf-8')
    render_docx(run.project,run.chapters,tmp_path/'docx')
    assert verification_report(run,tmp_path)['pipeline_complete']
    document=Document(tmp_path/'docx')
    paragraph=next(p for p in document.paragraphs if 'First line' in p.text)
    paragraph.text=paragraph.text.replace('\n','')
    document.save(tmp_path/'docx')
    assert not verification_report(run,tmp_path)['export_integrity']['docx']['matches_current_source']


def test_custom_story_keeps_brief_and_targets_but_routes_only_to_ollama(tmp_path):
    path = tmp_path / "story.json"
    path.write_text(json.dumps({"title": "A Different Book", "premise": "A botanist negotiates peace between rival gardens.",
        "requested_chapters": 24, "desired_word_count": 60000, "min_words_per_chapter": 2000,
        "max_words_per_chapter": 3000, "preferred_provider_name": "openai",
        "story_brief": {"ending_target": "Both gardens survive; the botanist leaves."},
        "task_routing": {"chapter_draft": {"provider_name": "openai", "model_name": "remote"}}}))
    project = load_project_spec(path, "local-model")
    assert project.preferred_provider_name == "ollama"
    assert project.preferred_model == "local-model"
    assert project.desired_word_count == 60000
    assert project.story_brief.ending_target == "Both gardens survive; the botanist leaves."
    assert not any(project.task_routing.model_dump().values())


def test_custom_story_rejects_unreachable_length_before_startup(tmp_path):
    path = tmp_path / "story.json"
    path.write_text(json.dumps({"title": "Story", "premise": "A botanist saves a garden.",
        "requested_chapters": 2, "desired_word_count": 60000, "min_words_per_chapter": 100, "max_words_per_chapter": 200}))
    with pytest.raises(ValueError, match="unreachable"):
        load_project_spec(path, "local-model")
