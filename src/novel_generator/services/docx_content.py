"""Render manuscript Markdown as editable Word content without fetching assets."""
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from markdown_it import MarkdownIt


def _inline(paragraph, tokens, *, header=False):
    bold = italic = strike = 0
    links = []
    for token in tokens:
        if token.type in {"strong_open", "strong_close"}:
            bold += token.nesting
        elif token.type in {"em_open", "em_close"}:
            italic += token.nesting
        elif token.type in {"s_open", "s_close"}:
            strike += token.nesting
        elif token.type in {"softbreak", "hardbreak"}:
            paragraph.add_run().add_break()
        elif token.type == "link_open":
            links.append((token.attrGet("href"), len(paragraph.text)))
        elif token.type == "link_close":
            href, start = links.pop()
            if href and paragraph.text[start:] != href:
                paragraph.add_run(f" ({href})")
        else:
            # Image captions and their reference paths remain visible. This
            # text-only exporter never downloads or invents an illustration.
            text = token.content
            if token.type == "image" and token.attrGet("src"):
                text += f" ({token.attrGet('src')})"
            run = paragraph.add_run(text)
            run.bold = True if bold or header else None
            run.italic = True if italic else None
            run.font.strike = True if strike else None
            if token.type == "code_inline":
                run.font.name = "Courier New"


def _table(document, tokens):
    rows, row, cell, header = [], None, None, False
    for token in tokens:
        if token.type == "tr_open":
            row = []
            rows.append(row)
        elif token.type in {"th_open", "td_open"}:
            header = token.type == "th_open"
            cell = []
            row.append((header, cell))
        elif token.type == "inline":
            cell.extend(token.children or [])
    if not rows:
        return
    table = document.add_table(rows=0, cols=max(len(row) for row in rows))
    table.style = "Table Grid"
    for row in rows:
        cells = table.add_row().cells
        for target, (header, children) in zip(cells, row):
            paragraph = target.paragraphs[0]
            paragraph.paragraph_format.first_line_indent = Inches(0)
            _inline(paragraph, children, header=header)


def add_manuscript_content(document, content):
    # CommonMark handles literal underscores, nested emphasis, escaped symbols,
    # and code correctly. Disable raw HTML; enable ordinary reference tables.
    tokens = MarkdownIt("commonmark", {"html": False}).enable(["table", "strikethrough"]).parse(content)
    lists, quote_depth, paragraph = [], 0, None
    index = 0
    while index < len(tokens):
        token = tokens[index]
        if token.type == "table_open":
            end = index + 1
            while tokens[end].type != "table_close":
                end += 1
            _table(document, tokens[index + 1:end])
            index = end
        elif token.type == "blockquote_open":
            quote_depth += 1
        elif token.type == "blockquote_close":
            quote_depth -= 1
        elif token.type in {"bullet_list_open", "ordered_list_open"}:
            start = token.attrGet("start")
            lists.append({"ordered": token.type == "ordered_list_open", "next": int(start) if start is not None else 1, "pending": False})
        elif token.type in {"bullet_list_close", "ordered_list_close"}:
            lists.pop()
        elif token.type == "list_item_open":
            lists[-1]["number"] = lists[-1]["next"]
            lists[-1]["next"] += 1
            lists[-1]["pending"] = True
        elif token.type == "heading_open":
            paragraph = document.add_heading("", level=min(9, int(token.tag[1:]) + 1))
        elif token.type == "paragraph_open":
            style = "Quote" if quote_depth else None
            prefix = ""
            if lists:
                current = lists[-1]
                style = "List Paragraph"
                if current["pending"]:
                    prefix = f"{current['number']}. " if current["ordered"] else "\u2022 "
                    current["pending"] = False
            paragraph = document.add_paragraph(style=style)
            if lists or quote_depth:
                paragraph.paragraph_format.first_line_indent = Inches(0)
                paragraph.paragraph_format.left_indent = Inches(0.25 * (len(lists) + quote_depth))
            if prefix:
                paragraph.add_run(prefix)
        elif token.type == "inline":
            _inline(paragraph, token.children or [])
        elif token.type in {"paragraph_close", "heading_close"}:
            paragraph = None
        elif token.type in {"fence", "code_block"}:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.first_line_indent = Inches(0)
            run = paragraph.add_run(token.content)
            run.font.name = "Courier New"
            run.font.size = Pt(10)
        elif token.type == "hr":
            paragraph = document.add_paragraph("\u2042")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        index += 1
